#!/usr/bin/env python3
"""Extract, chunk, hash, and index long-form research documents."""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import re
import shutil
import sqlite3
import sys
import zipfile
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree

from longform_common import (
    LongformError,
    append_jsonl,
    load_project,
    reexec_in_runtime,
    update_pipeline_stage,
)

SUPPORTED = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".tsv",
    ".json",
    ".html",
    ".htm",
    ".pdf",
    ".docx",
    ".pptx",
    ".srt",
    ".vtt",
    ".log",
}


class TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"script", "style", "noscript"}:
            self.skip_depth += 1
        if tag.lower() in {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "tr"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "noscript"} and self.skip_depth:
            self.skip_depth -= 1
        if tag.lower() in {"p", "div", "li", "h1", "h2", "h3", "h4", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.skip_depth:
            self.parts.append(data)

    def text(self) -> str:
        value = html.unescape("".join(self.parts))
        return normalize_text(value)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def safe_id(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9\u4e00-\u9fff._-]+", "-", value).strip(".-")
    return value[:80] or "source"


def normalize_text(value: str) -> str:
    value = value.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n[ \t]+", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def read_plain(path: Path) -> list[tuple[str, str]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return [("document", normalize_text(text))]


def read_json(path: Path) -> list[tuple[str, str]]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        text = json.dumps(payload, ensure_ascii=False, indent=2)
    except json.JSONDecodeError:
        text = path.read_text(encoding="utf-8", errors="replace")
    return [("document", normalize_text(text))]


def read_delimited(path: Path) -> list[tuple[str, str]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    rows: list[str] = []
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for row in reader:
            rows.append(" | ".join(cell.strip() for cell in row))
    return [("document", normalize_text("\n".join(rows)))]


def read_html(path: Path) -> list[tuple[str, str]]:
    parser = TextExtractor()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    return [("document", parser.text())]


def read_subtitles(path: Path) -> list[tuple[str, str]]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    lines: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped or stripped.isdigit() or stripped.upper() == "WEBVTT":
            continue
        if "-->" in stripped or stripped.startswith(("NOTE", "STYLE", "REGION")):
            continue
        stripped = re.sub(r"<[^>]+>", "", stripped)
        lines.append(stripped)
    return [("subtitle", normalize_text("\n".join(lines)))]


def read_docx(path: Path) -> list[tuple[str, str]]:
    reexec_in_runtime("docx")
    from docx import Document  # type: ignore

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [("document", normalize_text("\n".join(parts)))]


def read_pdf(path: Path) -> list[tuple[str, str]]:
    reexec_in_runtime("pypdf")
    from pypdf import PdfReader  # type: ignore

    reader = PdfReader(str(path))
    sections: list[tuple[str, str]] = []
    for number, page in enumerate(reader.pages, 1):
        text = normalize_text(page.extract_text() or "")
        if text:
            sections.append((f"page={number}", text))
    return sections


def slide_sort_key(name: str) -> tuple[int, str]:
    match = re.search(r"slide(\d+)\.xml$", name)
    return (int(match.group(1)) if match else 10**9, name)


def read_pptx(path: Path) -> list[tuple[str, str]]:
    sections: list[tuple[str, str]] = []
    with zipfile.ZipFile(path) as archive:
        slide_names = sorted(
            (
                name
                for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            ),
            key=slide_sort_key,
        )
        for index, name in enumerate(slide_names, 1):
            root = ElementTree.fromstring(archive.read(name))
            texts = [node.text or "" for node in root.iter() if node.tag.endswith("}t")]
            text = normalize_text("\n".join(texts))
            if text:
                sections.append((f"slide={index}", text))
    return sections


def extract(path: Path) -> tuple[str, list[tuple[str, str]]]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown", ".log"}:
        return "document", read_plain(path)
    if suffix in {".csv", ".tsv"}:
        return "document", read_delimited(path)
    if suffix == ".json":
        return "document", read_json(path)
    if suffix in {".html", ".htm"}:
        return "document", read_html(path)
    if suffix in {".srt", ".vtt"}:
        return "subtitle", read_subtitles(path)
    if suffix == ".docx":
        return "document", read_docx(path)
    if suffix == ".pdf":
        return "document", read_pdf(path)
    if suffix == ".pptx":
        return "document", read_pptx(path)
    raise LongformError(f"unsupported document type: {path}")


def collect_inputs(values: list[Path]) -> list[Path]:
    files: list[Path] = []
    for value in values:
        path = value.expanduser().resolve()
        if path.is_file() and path.suffix.lower() in SUPPORTED:
            files.append(path)
        elif path.is_dir():
            files.extend(
                candidate
                for candidate in path.rglob("*")
                if candidate.is_file() and candidate.suffix.lower() in SUPPORTED
            )
        else:
            print(f"Skipping unsupported or missing input: {path}", file=sys.stderr)
    return sorted(set(files))


def split_paragraph(paragraph: str, max_chars: int) -> list[str]:
    if len(paragraph) <= max_chars:
        return [paragraph]
    sentences = re.split(r"(?<=[。！？!?；;\.])\s*", paragraph)
    result: list[str] = []
    current = ""
    for sentence in sentences:
        sentence = sentence.strip()
        if not sentence:
            continue
        if len(sentence) > max_chars:
            if current:
                result.append(current)
                current = ""
            for start in range(0, len(sentence), max_chars):
                result.append(sentence[start : start + max_chars])
            continue
        candidate = sentence if not current else current + sentence
        if len(candidate) > max_chars:
            result.append(current)
            current = sentence
        else:
            current = candidate
    if current:
        result.append(current)
    return result


def chunk_sections(
    sections: list[tuple[str, str]], max_chars: int, overlap: int
) -> Iterable[tuple[str, str]]:
    for locator, text in sections:
        paragraphs = [part.strip() for part in text.split("\n\n") if part.strip()]
        current = ""
        chunk_number = 0
        for paragraph in paragraphs:
            for piece in split_paragraph(paragraph, max_chars):
                candidate = piece if not current else current + "\n\n" + piece
                if len(candidate) > max_chars and current:
                    chunk_number += 1
                    yield f"{locator};chunk={chunk_number}", current
                    tail = current[-overlap:] if overlap else ""
                    current = (tail + "\n\n" + piece).strip()
                else:
                    current = candidate
        if current:
            chunk_number += 1
            yield f"{locator};chunk={chunk_number}", current


def init_database(path: Path, reset: bool) -> tuple[sqlite3.Connection, bool]:
    if reset and path.exists():
        path.unlink()
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path)
    connection.execute(
        "CREATE TABLE IF NOT EXISTS sources ("
        "source_id TEXT PRIMARY KEY, path TEXT, sha256 TEXT, kind TEXT, title TEXT, "
        "characters INTEGER, extracted_at TEXT)"
    )
    try:
        connection.execute(
            "CREATE VIRTUAL TABLE IF NOT EXISTS chunks USING fts5("
            "chunk_id UNINDEXED, source_id UNINDEXED, citation UNINDEXED, text)"
        )
        fts = True
    except sqlite3.OperationalError:
        connection.execute(
            "CREATE TABLE IF NOT EXISTS chunks ("
            "chunk_id TEXT PRIMARY KEY, source_id TEXT, citation TEXT, text TEXT)"
        )
        fts = False
    return connection, fts


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project", type=Path, required=True)
    parser.add_argument(
        "--input",
        type=Path,
        action="append",
        help="File or directory. May be repeated. Defaults to research/inbox.",
    )
    parser.add_argument("--chunk-chars", type=int, default=2200)
    parser.add_argument("--overlap-chars", type=int, default=220)
    parser.add_argument("--no-copy", action="store_true")
    parser.add_argument("--reset-index", action="store_true")
    args = parser.parse_args()
    if not 500 <= args.chunk_chars <= 10000:
        parser.error("--chunk-chars must be between 500 and 10000")
    if not 0 <= args.overlap_chars < args.chunk_chars:
        parser.error("--overlap-chars must be smaller than --chunk-chars")
    return args


def main() -> int:
    args = parse_args()
    project, _ = load_project(args.project)
    inputs = args.input or [project / "research" / "inbox"]
    files = collect_inputs(inputs)
    if not files:
        raise LongformError("no supported research documents were found")

    database_path = project / "research" / "index.sqlite"
    connection, fts = init_database(database_path, args.reset_index)
    sources_jsonl = project / "research" / "sources.jsonl"
    chunks_jsonl = project / "research" / "chunks.jsonl"
    if args.reset_index:
        sources_jsonl.write_text("", encoding="utf-8")
        chunks_jsonl.write_text("", encoding="utf-8")
    existing_hashes = {
        row[0] for row in connection.execute("SELECT sha256 FROM sources").fetchall()
    }
    processed = 0
    chunk_count = 0
    skipped = 0
    update_pipeline_stage(project, "corpus", "running", files=len(files))

    try:
        for original in files:
            digest = sha256(original)
            if digest in existing_hashes:
                skipped += 1
                continue
            stored = original
            if not args.no_copy:
                try:
                    original.relative_to(project)
                except ValueError:
                    destination = project / "research" / "inbox" / original.name
                    if destination.exists() and sha256(destination) != digest:
                        destination = destination.with_name(
                            f"{destination.stem}-{digest[:8]}{destination.suffix}"
                        )
                    if not destination.exists():
                        shutil.copy2(original, destination)
                    stored = destination

            kind, sections = extract(stored)
            total_text = "\n\n".join(text for _, text in sections if text.strip())
            if not total_text.strip():
                print(f"No readable text: {stored}", file=sys.stderr)
                skipped += 1
                continue
            source_id = safe_id(f"{stored.stem}-{digest[:10]}")
            extracted_path = project / "research" / "extracted" / f"{source_id}.md"
            extracted_parts = [
                f"# {stored.name}",
                "",
                f"- Source ID: `{source_id}`",
                f"- SHA-256: `{digest}`",
                f"- Local path: `{stored}`",
                "",
            ]
            for locator, text in sections:
                extracted_parts.extend([f"## {locator}", "", text, ""])
            extracted_path.write_text("\n".join(extracted_parts), encoding="utf-8")
            extracted_at = datetime.now(timezone.utc).isoformat()
            source_entry = {
                "source_id": source_id,
                "title": stored.name,
                "kind": kind,
                "local_path": str(stored),
                "extracted_path": str(extracted_path.relative_to(project)),
                "sha256": digest,
                "characters": len(total_text),
                "sections": len(sections),
                "extracted_at": extracted_at,
            }
            append_jsonl(sources_jsonl, source_entry)
            connection.execute(
                "INSERT OR REPLACE INTO sources "
                "(source_id, path, sha256, kind, title, characters, extracted_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?)",
                (
                    source_id,
                    str(stored),
                    digest,
                    kind,
                    stored.name,
                    len(total_text),
                    extracted_at,
                ),
            )
            for index, (locator, text) in enumerate(
                chunk_sections(sections, args.chunk_chars, args.overlap_chars), 1
            ):
                chunk_id = f"{source_id}:{index:04d}"
                citation = f"{source_id}#{locator}"
                entry = {
                    "chunk_id": chunk_id,
                    "source_id": source_id,
                    "citation": citation,
                    "characters": len(text),
                    "text": text,
                }
                append_jsonl(chunks_jsonl, entry)
                if fts:
                    connection.execute("DELETE FROM chunks WHERE chunk_id = ?", (chunk_id,))
                    connection.execute(
                        "INSERT INTO chunks (chunk_id, source_id, citation, text) VALUES (?, ?, ?, ?)",
                        (chunk_id, source_id, citation, text),
                    )
                else:
                    connection.execute(
                        "INSERT OR REPLACE INTO chunks (chunk_id, source_id, citation, text) "
                        "VALUES (?, ?, ?, ?)",
                        (chunk_id, source_id, citation, text),
                    )
                chunk_count += 1
            connection.commit()
            existing_hashes.add(digest)
            processed += 1
            print(f"Indexed: {stored.name} -> {source_id}")
    except Exception:
        connection.rollback()
        update_pipeline_stage(project, "corpus", "failed")
        raise
    finally:
        connection.close()

    update_pipeline_stage(
        project,
        "corpus",
        "completed",
        processed=processed,
        chunks=chunk_count,
        skipped=skipped,
        fts5=fts,
    )
    print(
        f"Corpus ready: processed={processed}, chunks={chunk_count}, "
        f"skipped={skipped}, fts5={fts}"
    )
    print(f"Index: {database_path}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (LongformError, OSError, ValueError, zipfile.BadZipFile) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
